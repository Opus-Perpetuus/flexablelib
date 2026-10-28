"""
Excel import / export.

Round-trips plain tables and merged regions through openpyxl workbooks.
"""

from pathlib import Path

import pytest

from flexablelib import Table


pytest.importorskip('openpyxl')


def build_table():
    table = Table()
    table.add_column('Region', ['Norte', 'Norte', 'Sur', 'Sur'])
    table.add_column('Month', ['Jan', 'Feb', 'Jan', 'Feb'])
    table.add_column('Sales', [120.5, 98.25, 210.0, 175.75])
    return table


def test_to_excel_and_from_excel_round_trip_values(tmp_path):
    path = tmp_path / 'plain.xlsx'
    table = build_table()
    table.to_excel(path)

    loaded = Table.from_excel(path)

    assert loaded.headers == ['Region', 'Month', 'Sales']
    assert loaded.rows[0][0] == 'Norte'
    assert loaded.rows[0][2] == 120.5


def test_to_excel_writes_merged_cells(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / 'merged.xlsx'
    table = build_table()
    table.merge_cells(0, 0, last_row=1, value='Norte')
    table.merge_cells(0, 1, last_column=2, value='Q1')
    table.to_excel(path)

    workbook = load_workbook(path)
    ranges = {str(r) for r in workbook.active.merged_cells.ranges}
    workbook.close()

    assert 'A2:A3' in ranges
    assert 'B2:C2' in ranges


def test_from_excel_restores_merged_regions(tmp_path):
    path = tmp_path / 'round.xlsx'
    table = build_table()
    table.merge_cells(0, 0, last_row=1, value='Norte')
    table.merge_cells(2, 0, last_column=1, value='Sur block')
    table.to_excel(path)

    loaded = Table.from_excel(path)

    assert len(loaded.merged_regions) == 2
    regions = {(r.first_row, r.first_column, r.last_row, r.last_column)
               for r in loaded.merged_regions}
    assert (0, 0, 1, 0) in regions
    assert (2, 0, 2, 1) in regions
    assert 'Norte' in str(loaded)
    assert 'Sur block' in str(loaded)


def test_to_excel_shifts_merges_when_index_included(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / 'indexed.xlsx'
    table = build_table()
    table.show_index = True
    table.merge_cells(0, 0, last_row=1, value='Norte')
    table.to_excel(path, include_index=True)

    workbook = load_workbook(path)
    ranges = {str(r) for r in workbook.active.merged_cells.ranges}
    workbook.close()

    # Index column is A; body Region is B → B2:B3
    assert 'B2:B3' in ranges


def test_from_excel_sheet_selection(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / 'multi.xlsx'
    workbook = Workbook()
    first = workbook.active
    first.title = 'One'
    first.append(['A'])
    first.append([1])
    second = workbook.create_sheet('Two')
    second.append(['X', 'Y'])
    second.append(['x1', 'y1'])
    second.append(['x2', 'y2'])
    second.merge_cells('A2:A3')
    workbook.save(path)
    workbook.close()

    loaded = Table.from_excel(path, sheet='Two')
    assert loaded.headers == ['X', 'Y']
    assert len(loaded.merged_regions) == 1
    region = loaded.merged_regions[0]
    assert (region.first_row, region.first_column, region.last_row, region.last_column) == (0, 0, 1, 0)
